import os
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

def create_deck(output_paths):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Minimalist Palette
    COLOR_BG = RGBColor(248, 250, 252)          # Slate-50 (Crisp, clean light canvas)
    COLOR_CARD_BG = RGBColor(255, 255, 255)     # Pure White card
    COLOR_CARD_BORDER = RGBColor(226, 232, 240) # Slate-200 subtle border
    COLOR_TEXT_TITLE = RGBColor(15, 23, 42)     # Slate-900 bold title
    COLOR_TEXT_SUB = RGBColor(71, 85, 105)      # Slate-600 subtitle
    COLOR_TEXT_BODY = RGBColor(51, 65, 85)      # Slate-700 body
    COLOR_TEXT_MUTED = RGBColor(100, 116, 139)  # Slate-500 captions
    COLOR_PRIMARY = RGBColor(79, 70, 229)       # Indigo-600 (Restrained accent)
    COLOR_PRIMARY_BG = RGBColor(238, 242, 255)  # Indigo-50 pill bg
    COLOR_CHIP_BG = RGBColor(241, 245, 249)     # Slate-100 tag bg

    def apply_bg(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = COLOR_BG
        bg.line.fill.background() # No border
        return bg

    def add_header(slide, kicker_text, slide_num):
        # Kicker Pill
        pill = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.55), Inches(3.4), Inches(0.38))
        pill.fill.solid()
        pill.fill.fore_color.rgb = COLOR_PRIMARY_BG
        pill.line.color.rgb = RGBColor(199, 210, 254)
        pill.line.width = Pt(1)
        tf = pill.text_frame
        tf.word_wrap = False
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.text = kicker_text
        p.alignment = PP_ALIGN.CENTER
        p.font.size = Pt(10)
        p.font.bold = True
        p.font.color.rgb = COLOR_PRIMARY
        p.font.name = "Segoe UI"

        # Slide Number
        num_box = slide.shapes.add_textbox(Inches(11.5), Inches(0.55), Inches(1.0), Inches(0.38))
        tf_num = num_box.text_frame
        tf_num.vertical_anchor = MSO_ANCHOR.MIDDLE
        p_num = tf_num.paragraphs[0]
        p_num.text = slide_num
        p_num.alignment = PP_ALIGN.RIGHT
        p_num.font.size = Pt(11)
        p_num.font.color.rgb = COLOR_TEXT_MUTED
        p_num.font.name = "Segoe UI"

    def add_title_area(slide, title, subtitle):
        tb = slide.shapes.add_textbox(Inches(0.8), Inches(1.05), Inches(11.733), Inches(1.2))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        
        p_title = tf.paragraphs[0]
        p_title.text = title
        p_title.font.size = Pt(28)
        p_title.font.bold = True
        p_title.font.color.rgb = COLOR_TEXT_TITLE
        p_title.font.name = "Segoe UI"
        p_title.space_after = Pt(4)

        p_sub = tf.add_paragraph()
        p_sub.text = subtitle
        p_sub.font.size = Pt(14)
        p_sub.font.color.rgb = COLOR_TEXT_SUB
        p_sub.font.name = "Segoe UI"

    # =========================================================================
    # SLIDE 1: PROBLEM STATEMENT
    # =========================================================================
    slide1 = prs.slides.add_slide(blank_layout)
    apply_bg(slide1)
    add_header(slide1, "PRAGATI · HORIZON HOLLOW HACKATHON", "01 / 03")
    add_title_area(
        slide1,
        "The AI Tutor Amnesia Problem",
        "Why generic conversational models and flat quiz platforms fail modern STEM learners"
    )

    cards1 = [
        (
            "01. Zero Cross-Session Memory",
            "Standard chat models reset completely after every session. Students are forced to re-explain their curriculum, past mistakes, and weak spots from scratch every single day.",
            "Impact: 0% continuity across learning chats",
            RGBColor(225, 29, 72) # Rose-600
        ),
        (
            "02. Answer-Spoiling vs Mentoring",
            "Generic chatbots provide final code or answers immediately. By eliminating the productive struggle of problem-solving, cognitive retention collapses rapidly.",
            "Ebbinghaus Curve: 80% forgotten within 48 hours",
            RGBColor(217, 119, 6) # Amber-600
        ),
        (
            "03. Flat, One-Size-Fits-All Quizzes",
            "Traditional online quizzes present identical questions regardless of whether a student has 10% or 90% mastery. There is no dwell tracking, hint scoring, or difficulty adaptation.",
            "Result: Disengagement and false confidence",
            RGBColor(79, 70, 229) # Indigo-600
        )
    ]

    card_w = Inches(3.68)
    card_h = Inches(4.35)
    card_y = Inches(2.35)
    gap = Inches(0.34)

    for i, (ctitle, cdesc, ctag, tag_color) in enumerate(cards1):
        cx = Inches(0.8) + i * (card_w + gap)
        card = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, cx, card_y, card_w, card_h)
        card.fill.solid()
        card.fill.fore_color.rgb = COLOR_CARD_BG
        card.line.color.rgb = COLOR_CARD_BORDER
        card.line.width = Pt(1.2)

        # Top Accent stripe inside card
        stripe = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, cx + Inches(0.3), card_y + Inches(0.35), Inches(0.6), Inches(0.06))
        stripe.fill.solid()
        stripe.fill.fore_color.rgb = tag_color
        stripe.line.fill.background()

        # Text inside Card
        tb = slide1.shapes.add_textbox(cx + Inches(0.3), card_y + Inches(0.55), card_w - Inches(0.6), card_h - Inches(0.8))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        p1 = tf.paragraphs[0]
        p1.text = ctitle
        p1.font.size = Pt(17)
        p1.font.bold = True
        p1.font.color.rgb = COLOR_TEXT_TITLE
        p1.font.name = "Segoe UI"
        p1.space_after = Pt(12)

        p2 = tf.add_paragraph()
        p2.text = cdesc
        p2.font.size = Pt(12.5)
        p2.font.color.rgb = COLOR_TEXT_BODY
        p2.font.name = "Segoe UI"
        p2.space_after = Pt(20)

        p3 = tf.add_paragraph()
        p3.text = ctag
        p3.font.size = Pt(11)
        p3.font.bold = True
        p3.font.color.rgb = tag_color
        p3.font.name = "Segoe UI"

    # =========================================================================
    # SLIDE 2: THE SOLUTION
    # =========================================================================
    slide2 = prs.slides.add_slide(blank_layout)
    apply_bg(slide2)
    add_header(slide2, "PRAGATI · THE SOLUTION", "02 / 03")
    add_title_area(
        slide2,
        "Pragati: The Self-Evolving Learning Companion",
        "Persistent goal memory, autonomous mastery tooling, and dynamically calibrated difficulty"
    )

    cards2 = [
        (
            "Persistent Goal Memory",
            "Student goals (e.g. 'DSA', 'Calculus') and fine-grained subtopics persist in real-time database state. Every chat session loads complete mastery history, delivering tailored context.",
            [
                "Remembers topics & progress across all chats",
                "Subtopic mastery tracked to the millisecond",
                "Zero repetition needed from the student"
            ]
        ),
        (
            "Autonomous Mastery Tooling",
            "On clicking 'Test me', a clean topic prompt is sent. The AI agent executes check_topic_mastery in the background, inspects live mastery, and actively asks which subtopic to test.",
            [
                "Agent acts as active tutor, not vending machine",
                "Subtopic choice negotiated Socratically",
                "Zero prompt engineering required by user"
            ]
        ),
        (
            "4-Tier Difficulty Ladder",
            "Quizzes calibrate dynamically to the student's current proficiency level to ensure high engagement and optimal cognitive challenge:",
            [
                "Beginner (< 50% Mastery): Foundational drills",
                "Intermediate (50-79%): Application & analysis",
                "Advanced (80-89%): Edge cases & synthesis",
                "Expert (≥ 90%): Rigorous near-exam transfer"
            ]
        )
    ]

    for i, (ctitle, cdesc, bullets) in enumerate(cards2):
        cx = Inches(0.8) + i * (card_w + gap)
        card = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, cx, card_y, card_w, card_h)
        card.fill.solid()
        card.fill.fore_color.rgb = COLOR_CARD_BG
        card.line.color.rgb = COLOR_CARD_BORDER
        card.line.width = Pt(1.2)

        # Top Accent stripe inside card
        stripe = slide2.shapes.add_shape(MSO_SHAPE.RECTANGLE, cx + Inches(0.3), card_y + Inches(0.35), Inches(0.6), Inches(0.06))
        stripe.fill.solid()
        stripe.fill.fore_color.rgb = COLOR_PRIMARY
        stripe.line.fill.background()

        tb = slide2.shapes.add_textbox(cx + Inches(0.3), card_y + Inches(0.55), card_w - Inches(0.6), card_h - Inches(0.8))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        p1 = tf.paragraphs[0]
        p1.text = ctitle
        p1.font.size = Pt(17)
        p1.font.bold = True
        p1.font.color.rgb = COLOR_TEXT_TITLE
        p1.font.name = "Segoe UI"
        p1.space_after = Pt(10)

        p2 = tf.add_paragraph()
        p2.text = cdesc
        p2.font.size = Pt(12)
        p2.font.color.rgb = COLOR_TEXT_BODY
        p2.font.name = "Segoe UI"
        p2.space_after = Pt(12)

        for b in bullets:
            pb = tf.add_paragraph()
            pb.text = f"•  {b}"
            pb.font.size = Pt(11)
            pb.font.color.rgb = COLOR_TEXT_SUB
            pb.font.name = "Segoe UI"
            pb.space_after = Pt(4)

    # =========================================================================
    # SLIDE 3: ARCHITECTURE & DEMO TRANSITION
    # =========================================================================
    slide3 = prs.slides.add_slide(blank_layout)
    apply_bg(slide3)
    add_header(slide3, "PRAGATI · ARCHITECTURE & DEMONSTRATION", "03 / 03")
    add_title_area(
        slide3,
        "Engineered for Production Rigor",
        "Clean schemas, zero answer leakage, and live telemetry synchronization"
    )

    col1_w = Inches(6.8)
    col2_w = Inches(4.6)
    col_gap = Inches(0.33)

    # Left Card: Technical Highlights
    card_l = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), card_y, col1_w, card_h)
    card_l.fill.solid()
    card_l.fill.fore_color.rgb = COLOR_CARD_BG
    card_l.line.color.rgb = COLOR_CARD_BORDER
    card_l.line.width = Pt(1.2)

    tb_l = slide3.shapes.add_textbox(Inches(1.15), card_y + Inches(0.35), col1_w - Inches(0.7), card_h - Inches(0.7))
    tf_l = tb_l.text_frame
    tf_l.word_wrap = True
    tf_l.margin_left = tf_l.margin_top = tf_l.margin_right = tf_l.margin_bottom = 0

    p_lh = tf_l.paragraphs[0]
    p_lh.text = "Production Architectural Pillars"
    p_lh.font.size = Pt(17)
    p_lh.font.bold = True
    p_lh.font.color.rgb = COLOR_TEXT_TITLE
    p_lh.font.name = "Segoe UI"
    p_lh.space_after = Pt(12)

    tech_points = [
        ("Dual Backend Parity", "Express.js and Supabase Edge Functions stay behaviorally identical, verified by 141 automated vitest unit and evaluation tests."),
        ("Zero Answer Leakage", "Server sanitization strips answer keys and question previews from LLM chat prose; interactive cards auto-render in the UI."),
        ("Multi-Factor Telemetry", "Granular per-question dwell timing, hint penalty deductions, and dynamic ELO skill adjustment."),
        ("Socratic Review Loop", "Missed questions populate Analytics 'Questions to Review', where the AI re-tutors until mastery is proven.")
    ]

    for title, desc in tech_points:
        pt = tf_l.add_paragraph()
        pt.text = f"{title}: "
        pt.font.bold = True
        pt.font.size = Pt(11.5)
        pt.font.color.rgb = COLOR_PRIMARY
        pt.font.name = "Segoe UI"
        
        # Add desc in same paragraph
        run = pt.add_run()
        run.text = desc
        run.font.bold = False
        run.font.size = Pt(11)
        run.font.color.rgb = COLOR_TEXT_BODY
        run.font.name = "Segoe UI"
        pt.space_after = Pt(8)

    # Right Card: Demo Walkthrough Plan
    card_r = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8) + col1_w + col_gap, card_y, col2_w, card_h)
    card_r.fill.solid()
    card_r.fill.fore_color.rgb = COLOR_PRIMARY_BG
    card_r.line.color.rgb = RGBColor(199, 210, 254)
    card_r.line.width = Pt(1.2)

    tb_r = slide3.shapes.add_textbox(Inches(0.8) + col1_w + col_gap + Inches(0.35), card_y + Inches(0.35), col2_w - Inches(0.7), card_h - Inches(0.7))
    tf_r = tb_r.text_frame
    tf_r.word_wrap = True
    tf_r.margin_left = tf_r.margin_top = tf_r.margin_right = tf_r.margin_bottom = 0

    p_rh = tf_r.paragraphs[0]
    p_rh.text = "Live Demonstration Plan"
    p_rh.font.size = Pt(17)
    p_rh.font.bold = True
    p_rh.font.color.rgb = COLOR_PRIMARY
    p_rh.font.name = "Segoe UI"
    p_rh.space_after = Pt(14)

    steps = [
        "1. Compact Workspace (3-chat limit, clean profile)",
        "2. 'My Topics' → Click 'Test me' on DSA goal",
        "3. Agent runs check_topic_mastery tool autonomously",
        "4. Socratic subtopic negotiation & Beginner tier",
        "5. Interactive Quiz Arena with LaTeX & timer",
        "6. Live mastery update & Analytics pipeline"
    ]

    for s in steps:
        ps = tf_r.add_paragraph()
        ps.text = s
        ps.font.size = Pt(11.5)
        ps.font.bold = False
        ps.font.color.rgb = COLOR_TEXT_TITLE
        ps.font.name = "Segoe UI"
        ps.space_after = Pt(8)

    # Save to all requested paths
    for path in output_paths:
        os.makedirs(os.path.dirname(os.path.abspath(path)), exist_ok=True)
        prs.save(path)
        print(f"Saved PPTX to: {path}")

if __name__ == '__main__':
    paths = [
        os.path.abspath("Pragati_Pitch_Deck.pptx"),
        os.path.abspath("client/public/Pragati_Pitch_Deck.pptx")
    ]
    create_deck(paths)
